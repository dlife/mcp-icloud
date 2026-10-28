"""Tekst uit mailbijlagen halen (PDF, Word, Excel, PowerPoint, platte tekst). Werkt op bytes, los van IMAP."""
import io
import os
import zipfile
from html.parser import HTMLParser

MAX_UNPACKED_BYTES = 200 * 1024 * 1024  # bescherming tegen zip-bommen in Office-bestanden

TEXT_TYPES = {"text/plain", "text/csv", "text/markdown", "text/calendar", "application/json", "application/xml", "text/xml"}
TEXT_EXTENSIONS = {".txt", ".csv", ".md", ".json", ".xml", ".ics", ".log", ".eml"}
OLD_OFFICE = {".doc", ".xls", ".ppt"}


class _HTMLText(HTMLParser):
    def __init__(self):
        super().__init__()
        self.parts: list[str] = []
        self._skip = 0

    def handle_starttag(self, tag, attrs):
        if tag in ("script", "style"):
            self._skip += 1
        elif tag in ("br", "p", "div", "tr", "li", "h1", "h2", "h3"):
            self.parts.append("\n")

    def handle_endtag(self, tag):
        if tag in ("script", "style") and self._skip:
            self._skip -= 1

    def handle_data(self, data):
        if not self._skip:
            self.parts.append(data)


def _decode(data: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", "replace")


def _check_zip(data: bytes) -> None:
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        if sum(i.file_size for i in zf.infolist()) > MAX_UNPACKED_BYTES:
            raise ValueError("Bestand is uitgepakt te groot")


def _pdf(data: bytes, max_chars: int) -> str:
    import pymupdf

    with pymupdf.open(stream=data, filetype="pdf") as doc:
        if doc.needs_pass:
            raise ValueError("PDF is beveiligd met een wachtwoord")
        out, total = [], 0
        for number, page in enumerate(doc, 1):
            text = page.get_text().strip()
            if text:
                out.append(f"--- pagina {number} ---\n{text}")
                total += len(text)
            if total > max_chars:
                break
        return "\n\n".join(out)


def _docx(data: bytes, max_chars: int) -> str:
    import docx

    _check_zip(data)
    document = docx.Document(io.BytesIO(data))
    lines = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            lines.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(lines)


def _xlsx(data: bytes, max_chars: int) -> str:
    import openpyxl

    _check_zip(data)
    book = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    out, total = [], 0
    for sheet in book.worksheets:
        out.append(f"--- blad {sheet.title} ---")
        for row in sheet.iter_rows(values_only=True):
            if not any(c is not None for c in row):
                continue
            line = " | ".join("" if c is None else str(c) for c in row)
            out.append(line)
            total += len(line)
            if total > max_chars:
                return "\n".join(out)
    return "\n".join(out)


def _pptx(data: bytes, max_chars: int) -> str:
    import pptx

    _check_zip(data)
    deck = pptx.Presentation(io.BytesIO(data))
    out = []
    for number, slide in enumerate(deck.slides, 1):
        out.append(f"--- dia {number} ---")
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                out.append(shape.text_frame.text.strip())
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    out.append(" | ".join(cell.text.strip() for cell in row.cells))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            out.append("Notities: " + slide.notes_slide.notes_text_frame.text.strip())
    return "\n".join(out)


def extract_text(name: str, content_type: str, data: bytes, max_chars: int = 50_000) -> dict:
    """Geef {name, type, size, text, truncated, note}. `note` legt uit waarom er geen (volledige) tekst is."""
    ext = os.path.splitext(name.lower())[1]
    ctype = content_type.lower()
    result = {"name": name, "type": ctype, "size": len(data), "text": "", "truncated": False, "note": ""}

    try:
        if ctype == "application/pdf" or ext == ".pdf":
            text = _pdf(data, max_chars)
            if not text.strip():
                result["note"] = "Geen tekst gevonden: waarschijnlijk een scan (OCR is nog niet beschikbaar)."
        elif ext == ".docx" or ctype.endswith("wordprocessingml.document"):
            text = _docx(data, max_chars)
        elif ext == ".xlsx" or ctype.endswith("spreadsheetml.sheet"):
            text = _xlsx(data, max_chars)
        elif ext == ".pptx" or ctype.endswith("presentationml.presentation"):
            text = _pptx(data, max_chars)
        elif ctype == "text/html" or ext in (".html", ".htm"):
            parser = _HTMLText()
            parser.feed(_decode(data))
            text = "".join(parser.parts)
        elif ctype in TEXT_TYPES or ext in TEXT_EXTENSIONS:
            text = _decode(data)
        elif ctype.startswith("image/"):
            text = ""
            result["note"] = "Afbeelding: tekstherkenning (OCR) is nog niet beschikbaar."
        elif ext in OLD_OFFICE:
            text = ""
            result["note"] = "Oud Office-formaat (.doc/.xls/.ppt) wordt niet ondersteund; vraag een .docx/.xlsx/.pptx of PDF."
        else:
            text = ""
            result["note"] = f"Bestandstype {ctype or ext or 'onbekend'} wordt niet ondersteund."
    except ValueError as exc:
        text, result["note"] = "", str(exc)
    except Exception as exc:  # kapot of vreemd bestand: geen details van de inhoud lekken
        text, result["note"] = "", f"Bestand kon niet gelezen worden ({type(exc).__name__})."

    text = text.strip()
    if len(text) > max_chars:
        text, result["truncated"] = text[:max_chars], True
    result["text"] = text
    return result
