import io

import docx
import pymupdf
import openpyxl
import pptx
import pytest

from icloud_mcp.attachments import extract_text


def make_pdf(*pages: str) -> bytes:
    doc = pymupdf.open()
    for text in pages:
        page = doc.new_page()
        if text:
            page.insert_text((72, 72), text)
    return doc.tobytes()


def make_docx() -> bytes:
    d = docx.Document()
    d.add_paragraph("Offerte voor Dieter")
    table = d.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Totaal"
    table.rows[0].cells[1].text = "125 euro"
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def make_xlsx() -> bytes:
    wb = openpyxl.Workbook()
    wb.active.title = "Kosten"
    wb.active.append(["Post", "Bedrag"])
    wb.active.append(["Huur", 900])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def make_pptx() -> bytes:
    deck = pptx.Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = "Planning"
    slide.placeholders[1].text = "Start in maart"
    buf = io.BytesIO()
    deck.save(buf)
    return buf.getvalue()


def test_pdf_text_with_page_markers():
    out = extract_text("f.pdf", "application/pdf", make_pdf("Factuur 2026-001", "Tweede pagina"))
    assert "Factuur 2026-001" in out["text"] and "pagina 2" in out["text"]
    assert out["note"] == ""


def test_scanned_pdf_gives_note_not_empty_silence():
    out = extract_text("scan.pdf", "application/pdf", make_pdf(""))
    assert out["text"] == "" and "scan" in out["note"]


def test_docx_includes_tables():
    out = extract_text("o.docx", "application/octet-stream", make_docx())
    assert "Offerte voor Dieter" in out["text"] and "Totaal | 125 euro" in out["text"]


def test_xlsx_rows_per_sheet():
    out = extract_text("k.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", make_xlsx())
    assert "blad Kosten" in out["text"] and "Huur | 900" in out["text"]


def test_pptx_slides():
    out = extract_text("p.pptx", "application/octet-stream", make_pptx())
    assert "Planning" in out["text"] and "Start in maart" in out["text"]


def test_plain_text_falls_back_to_cp1252():
    out = extract_text("n.txt", "text/plain", "café".encode("cp1252"))
    assert out["text"] == "café"


def test_html_is_stripped():
    out = extract_text("m.html", "text/html", b"<html><style>x{}</style><p>Hallo</p><script>evil()</script></html>")
    assert out["text"] == "Hallo"


def test_truncation_flag():
    out = extract_text("n.txt", "text/plain", b"a" * 100, max_chars=10)
    assert out["text"] == "a" * 10 and out["truncated"] is True


@pytest.mark.parametrize(
    "name,ctype,fragment",
    [("foto.jpg", "image/jpeg", "OCR"), ("oud.doc", "application/msword", "Oud Office"), ("x.bin", "application/octet-stream", "niet ondersteund")],
)
def test_unsupported_types_explain_why(name, ctype, fragment):
    out = extract_text(name, ctype, b"xxxx")
    assert out["text"] == "" and fragment in out["note"]


def test_corrupt_file_does_not_raise_or_leak():
    out = extract_text("kapot.pdf", "application/pdf", b"dit is geen pdf")
    assert out["text"] == "" and "kon niet gelezen worden" in out["note"]
